"""
Calibrate Mid-Sem 1 Presentation & Report to reserve combined GIS spatial mapping for Mid-Sem 2.
Keeps live Streamlit Diagnostic Console (Tab 1) for Mid-Sem 1 demo.
"""

from pathlib import Path
import docx
from pptx import Presentation

PROJECT_DIR = Path("/Users/nakul/Desktop/soil project")
PPTX_PATH = PROJECT_DIR / "Mid_Sem_Presentation_CSPC401.pptx"
DOCX_PATH = PROJECT_DIR / "Mid_Sem_Project_Report_CSPC401.docx"


def update_presentation():
    print(f"[*] Updating Presentation: {PPTX_PATH}")
    prs = Presentation(str(PPTX_PATH))

    # Slide 2: Outline
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "05" in p.text:
                    pass
                if "GIS Mapping, Web App & Future Roadmap" in p.text:
                    p.text = "Diagnostic Console, Phase 2 GIS Architecture & Project Roadmap"
                if "Kurukshetra Plot A orthomosaic zoning, live Streamlit Web App, and Phase 2 timeline." in p.text:
                    p.text = "Operational live diagnostic interface, planned Phase 2 GIS spatial mapping design, and 4-phase project roadmap."

    # Slide 10: Reframe GIS Mapping as Phase 2 / Mid-Sem 2 Architecture
    s10 = prs.slides[9]
    for shape in s10.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "4. Progress Achieved: Digital GIS Soil Health Mapping" in p.text:
                    p.text = "5. Phase 2 Architecture: Proposed Digital GIS Spatial Mapping (Mid-Sem 2)"
                if "AERIAL ORTHOMOSAIC MAPPING" in p.text:
                    p.text = "PROPOSED GIS SPATIAL TILING PIPELINE (SCHEDULED FOR MID-SEM 2)"
                if "Field Site: Kurukshetra Agricultural Plot A" in p.text:
                    p.text = "• Target Field Site: NIT Kurukshetra Agricultural Research Plots — Indo-Gangetic Alluvial testbed."
                if "Spatial Grid Tessellation: The large aerial orthomosaic is sliced into georeferenced spatial tiles (128x128 patches)" in p.text:
                    p.text = "• Spatial Grid Tessellation: Methodology to slice large aerial orthomosaics into sliding-window spatial tiles (224x224 px) representing contiguous GPS ground parcels."
                if "Zoning Classification: Each grid cell is scored for Soil Type, Health Status, and 0-100 SHI Index." in p.text:
                    p.text = "• Multi-Task Spatial Scoring: Each grid tile will be evaluated by UAV-SoilHealthNet to produce 3-color zoning (Healthy, Moderate, Degraded) and continuous SHI maps."
                if "Dual Deliverable: Produces high-res georeferenced PNG heatmaps + an interactive Leaflet HTML web map" in p.text:
                    p.text = "• Mid-Sem 2 Scheduled Deliverable: Demonstration of whole-field spatial GIS tiling, acreage statistics, and interactive clickable grid dashboards across diverse agro-climatic zones."

    # Slide 11: Live Diagnostic Console (Streamlit Tab 1) for Mid-Sem 1 Demo
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "4. Progress Achieved: Farmer Advisory & Streamlit Web App" in p.text:
                    p.text = "4. Progress Achieved: Live Diagnostic Console & Farmer Advisory (Mid-Sem 1 Demo)"
                if "💻 INTERACTIVE STREAMLIT WEB APP" in p.text:
                    p.text = "💻 OPERATIONAL STREAMLIT DIAGNOSTIC CONSOLE (TAB 1 LIVE DEMO)"
                if "GIS Drone Map Viewer: Pan, zoom, and inspect continuous soil health zoning of Kurukshetra agricultural plots." in p.text:
                    p.text = "✔ Automated Vegetation Removal: Built-in spectral filter eliminates weed canopy and inpaints bare soil texture before ML inference."
                if "Single Soil Diagnosis: Upload any high-res UAV soil tile and receive instant multi-task classification and SHI score." in p.text:
                    p.text = "✔ Real-Time Inference: Instant prediction of Soil Type, 3-class Degradation Status, and 0-100 SHI Score on real field photos."
                if "Command to Run: streamlit run web_app/app.py (Ready for live committee demo!)" in p.text:
                    p.text = "✔ Live Evaluation Demo Ready: streamlit run web_app/app.py (Diagnostic Console operational for committee review)."

    # Slide 12: Roadmap with clean 4-phase milestone split
    s12 = prs.slides[11]
    for shape in s12.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Phase 1: Foundation" in p.text:
                    p.text = "Phase 1: Foundation & Diagnostic AI (COMPLETED - Mid-Sem 1)  [Aug – Oct 2026]"
                if "Established baselines (~75-86%) and constructed initial UAV-SoilHealthNet prototype" in p.text:
                    p.text = "• Extracted 60 handcrafted agronomic features & established classical baselines (SVM 90.9%, RF 89.7%).\n• Engineered UAV-SoilHealthNet prototype (~80.6% Health Acc, 0.934 AUC) and Grad-CAM interpretability.\n• Deployed operational Streamlit Diagnostic Console with automated vegetation inpainting."
                if "Phase 2: Model Optimization & Drone Flights" in p.text:
                    p.text = "Phase 2: Combined GIS Spatial Mapping & Field Tiling (SCHEDULED - Mid-Sem 2)  [Nov 2026]"
                if "Fine-tune UAV-SoilHealthNet multi-task loss weighting" in p.text:
                    p.text = "• Deploy sliding-window orthomosaic grid tiler and GIS spatial zoning across 4 major Indian agro-zones.\n• Integrate interactive clickable Leaflet/HTML spatial dashboard and acreage distribution analytics.\n• Prepare edge quantization (ONNX/TensorRT) for UAV companion computers."
                if "Phase 3: Multimodal Calibration & WebODM Orthomosaics" in p.text:
                    p.text = "Phase 3: Autonomous Field Flights & Ground Sampling (SCHEDULED)  [Dec 2026 – Jan 2027]"
                if "Stitch raw drone frames into high-precision orthomosaics" in p.text:
                    p.text = "• Execute autonomous grid flights over NIT Kurukshetra campus plots at 15-30m AGL.\n• Collect 30-50 georeferenced physical soil core samples with GPS tags before rabi sowing.\n• Measure ground-truth laboratory pH, organic carbon (Walkley-Black), and moisture."
                if "Phase 4: Field Validation & Final Thesis Submission" in p.text:
                    p.text = "Phase 4: Model Optimization & Final Thesis Submission (SCHEDULED - End-Sem)  [Feb – Apr 2027]"
                if "Validate recommendations with local farming cooperatives" in p.text:
                    p.text = "• Fine-tune multi-task loss weighting to scale combined accuracy from ~80% to >92%.\n• Validate recommendations with local farming cooperatives in Haryana.\n• Final B.Tech thesis defense and research paper submission."

    # Slide 13: Summary
    s13 = prs.slides[12]
    for shape in s13.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Operational GIS & Web App: Demonstrated full spatial tiling on Kurukshetra Plot A" in p.text:
                    p.text = "✔  Operational Diagnostic Console & Clear Roadmap: Demonstrated live single-patch diagnosis and Grad-CAM explainability in Streamlit (Tab 1), with full GIS spatial mapping reserved for Mid-Sem 2 and drone flights for End-Sem."

    prs.save(str(PPTX_PATH))
    print("[+] Mid_Sem_Presentation_CSPC401.pptx successfully calibrated.")


def update_report():
    print(f"[*] Updating Report: {DOCX_PATH}")
    doc = docx.Document(str(DOCX_PATH))

    # Table 3: Milestone Timeline Table
    t3 = doc.tables[3]
    calibrated_milestones = [
        ["Phase 1 (Mid-Sem 1)", "UAV-SoilHealthNet Prototype, 60 Agronomic Features, Baseline Benchmarking & Streamlit Diagnostic Console", "Aug - Oct 2026", "COMPLETED ✔"],
        ["Phase 2 (Mid-Sem 2)", "Combined Digital GIS Spatial Field Mapping, Orthomosaic Grid Tiling (4 Agro-Zones) & Interactive Dashboard", "Nov 2026", "Scheduled"],
        ["Phase 3 (Pre-Sub)", "Autonomous UAV Flights over Kurukshetra plots & Physical GNSS Core Soil Sampling", "Dec 2026 - Jan 2027", "Scheduled"],
        ["Phase 4 (End-Sem)", "Laboratory Ground-Truth Assays (SOC, pH), Model Optimization (>92%) & Final Thesis Defense", "Feb - Apr 2027", "Scheduled"]
    ]

    while len(t3.rows) > 1:
        tr = t3.rows[-1]._tr
        t3._tbl.remove(tr)

    for row_data in calibrated_milestones:
        row = t3.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = val

    # Update discussion in text
    for p in doc.paragraphs:
        if "4.4 Digital GIS Soil Health Mapping" in p.text:
            p.text = "4.4 Digital GIS Spatial Field Mapping Architecture (Design for Mid-Sem 2)"
        if "Figure 4.4: Digital Soil Health GIS Zoning Map of Kurukshetra Agricultural Plot A" in p.text:
            p.text = "Figure 4.4: Proposed Digital Soil Health GIS Zoning Architecture (Scheduled for Mid-Semester 2 Demonstration)."
        if "demonstrates full spatial tiling on Kurukshetra Plot A with static and interactive HTML GIS maps, backed by an operational Streamlit GUI" in p.text:
            p.text = "delivers the core AI classification and explainability pipeline with an operational Streamlit diagnostic console (Tab 1). The combined whole-field GIS spatial mapping engine is scheduled as the primary deliverable for Mid-Semester 2, followed by physical UAV flights and model optimization in Phase 3 and 4"

    doc.save(str(DOCX_PATH))
    print("[+] Mid_Sem_Project_Report_CSPC401.docx successfully calibrated.")


if __name__ == "__main__":
    update_presentation()
    update_report()
